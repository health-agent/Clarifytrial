"""Attach editable chart-owned XLSX data to the two newly authored charts.

These worksheets contain the complete literal data used to create the poster.
They are not recovered source workbooks. No scientific result is changed.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path,PurePosixPath
import posixpath
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile,ZIP_DEFLATED
from openpyxl import Workbook

C='http://schemas.openxmlformats.org/drawingml/2006/chart'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
ET.register_namespace('c',C);ET.register_namespace('r',R)


def embed(source:Path,layout_path:Path,output:Path):
    if source.resolve()==output.resolve():raise ValueError('Preserve the authored source')
    charts=[e for e in json.loads(layout_path.read_text(encoding='utf-8'))['elements'] if e['kind']=='chart']
    with ZipFile(source) as z:
        parts={p:z.read(p) for p in z.namelist()}
    chart_paths=sorted(p for p in parts if re.search(r'/charts/chart\d+\.xml$',p))
    if len(charts)!=2 or len(chart_paths)!=2:raise ValueError('Exactly two poster charts expected')
    types=ET.fromstring(parts['[Content_Types].xml'])
    for number,(chart_path,expected) in enumerate(zip(chart_paths,charts),1):
        root=ET.fromstring(parts[chart_path]);series=root.findall(f'.//{{{C}}}ser')
        if len(series)!=2 or root.find(f'{{{C}}}externalData') is not None:raise ValueError('Expected new literal chart')
        wb=Workbook();ws=wb.active;ws.title='ChartData';ws['A1']='대상'
        for j,ser in enumerate(series):
            label=ser.find(f'{{{C}}}tx/{{{C}}}v')
            categories=ser.find(f'{{{C}}}cat/{{{C}}}strLit')
            values=ser.find(f'{{{C}}}val/{{{C}}}numLit')
            if label is None or categories is None or values is None:raise ValueError('Complete literal data required')
            values_found=[float(p.find(f'{{{C}}}v').text) for p in values.findall(f'{{{C}}}pt')]
            if len(values_found)!=1 or abs(values_found[0]-expected['values'][j])>1e-9:raise ValueError('Chart values differ from fixed results')
            cat_values=[p.find(f'{{{C}}}v').text or '' for p in categories.findall(f'{{{C}}}pt')]
            if len(cat_values)!=1:raise ValueError('One quantitative category expected')
            column=chr(ord('B')+j);ws.cell(1,j+2,label.text);ws.cell(2,1,cat_values[0]);ws.cell(2,j+2,values_found[0])
            tx=ser.find(f'{{{C}}}tx');tx.remove(label)
            ref=ET.SubElement(tx,f'{{{C}}}strRef');ET.SubElement(ref,f'{{{C}}}f').text=f'ChartData!${column}$1'
            cache=ET.SubElement(ref,f'{{{C}}}strCache');ET.SubElement(cache,f'{{{C}}}ptCount',{'val':'1'})
            pt=ET.SubElement(cache,f'{{{C}}}pt',{'idx':'0'});ET.SubElement(pt,f'{{{C}}}v').text=label.text
            for parent_name,literal,new_kind,cache_name,formula in (
                ('cat',categories,'strRef','strCache','ChartData!$A$2:$A$2'),
                ('val',values,'numRef','numCache',f'ChartData!${column}$2:${column}$2')):
                parent=ser.find(f'{{{C}}}{parent_name}');parent.remove(literal)
                reference=ET.SubElement(parent,f'{{{C}}}{new_kind}');ET.SubElement(reference,f'{{{C}}}f').text=formula
                copied=deepcopy(literal);copied.tag=f'{{{C}}}{cache_name}';reference.append(copied)
        stream=BytesIO();wb.save(stream)
        workbook_path=f'ppt/embeddings/clarifytrial-chart-{number}.xlsx';parts[workbook_path]=stream.getvalue()
        parent=PurePosixPath(chart_path).parent
        rel_path=str(parent/'_rels'/(PurePosixPath(chart_path).name+'.rels'))
        rels=ET.fromstring(parts[rel_path]) if rel_path in parts else ET.Element(f'{{{REL}}}Relationships')
        rid='rIdPosterWorkbook'
        if any(r.get('Id')==rid for r in rels):raise ValueError('Workbook relationship already exists')
        ET.SubElement(rels,f'{{{REL}}}Relationship',{'Id':rid,'Type':f'{R}/package','Target':posixpath.relpath(workbook_path,str(parent))})
        ET.register_namespace('',REL)
        parts[rel_path]=ET.tostring(rels,encoding='utf-8',xml_declaration=True)
        external=ET.SubElement(root,f'{{{C}}}externalData',{f'{{{R}}}id':rid});ET.SubElement(external,f'{{{C}}}autoUpdate',{'val':'0'})
        parts[chart_path]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
        ET.SubElement(types,f'{{{CT}}}Override',{'PartName':'/'+workbook_path,'ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
    ET.register_namespace('',CT)
    parts['[Content_Types].xml']=ET.tostring(types,encoding='utf-8',xml_declaration=True)
    output.parent.mkdir(parents=True,exist_ok=True)
    with ZipFile(output,'w',ZIP_DEFLATED) as z:
        for name,value in parts.items():z.writestr(name,value)
    print(f'Two chart-owned XLSX datasets embedded: {output}')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('source',type=Path);p.add_argument('layout',type=Path);p.add_argument('output',type=Path)
    a=p.parse_args();embed(a.source,a.layout,a.output)
